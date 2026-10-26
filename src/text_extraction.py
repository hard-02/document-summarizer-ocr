"""Extract raw text from uploaded documents (.txt, .pdf, .docx)."""
from __future__ import annotations

from pathlib import Path


def extract_text_from_txt(path: Path) -> str:
    """Read a plain-text file, tolerating non-UTF-8 encodings."""
    return Path(path).read_text(encoding="utf-8", errors="replace")


def extract_text_from_pdf(path: Path) -> str:
    """Extract text from every page of a PDF using PyMuPDF (fitz)."""
    import fitz  # noqa: PLC0415 (PyMuPDF)

    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def extract_text_from_docx(path: Path) -> str:
    """Extract paragraphs and table cells from a Word document."""
    import docx  # noqa: PLC0415 (python-docx)

    document = docx.Document(str(path))
    chunks = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    chunks.append(cell.text.strip())

    return "\n".join(chunks)


def extract_text(path: Path) -> str:
    """Dispatch to the right extractor based on the file extension."""
    suffix = Path(path).suffix.lower()
    if suffix == ".txt":
        return extract_text_from_txt(path)
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix == ".docx":
        return extract_text_from_docx(path)
    raise ValueError(f"Unsupported document type: {suffix!r} (use .txt, .pdf, or .docx)")
